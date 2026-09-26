"""
fetch_fjernvarme_mix.py - Fjernvarmens brændselsmix pr. kommune.

KONTEKST under Energi-dimensionen (vises, men indgår IKKE i scoren).
Viser hvor "ren" en kommunes fjernvarme reelt er: andel biomasse, affald,
fossilt og reelt vedvarende (sol/varmepumpe/overskudsvarme). Bevidst ikke
scoret, fordi næsten al dansk fjernvarme er afbrænding - en absolut renheds-
score ville gøre alle kommuner røde, og at score biomasse som "grønt" er et
omdiskuteret værdivalg (se metode-siden).

Kilde:
  Energistyrelsen, Energiproducenttællingen (EPT)
  "Produktion og brændselsforbrug per anlæg 2022-2024"
  https://ens.dk/media/7199/download  (.xlsx, ~2,2 MB)

Metode:
  - Filtrér til seneste år (2024) og varmeproducerende anlæg (varmeprod_TJ > 0).
  - Vægt hvert anlægs brændselsforbrug med andel_varmelev, så kun den del af
    brændslet der går til VARME tælles med (kraftvarmeværkers el-del fjernes).
  - Aggregér pr. kommune (vaerk_kommune) og beregn fire andele der summer til 100%:
      biomasse  = skovflis + træpiller + halm + træ/biomasseaffald + bio-olie + biogas
      affald    = affald
      fossil    = kul + fuelolie + spildolie + gasolie + raffinaderigas + LPG + naturgas
      ren       = solenergi + omgivelsesvarme (varmepumper) + elektricitet + brændselsfrit
  - Kommuner med under MIN_TJ samlet varme-input har ingen meningsfuld egen
    produktion: den er typisk kun en reservekedel (fx hovedstadskommuner på VEKS/
    CTR-nettet). De får i stedet mixet for det net, der forsyner dem (status
    "net"), fra ENS' "Fjernvarmenet 2022-2024" (https://ens.dk/media/8519/download,
    arket "125% metode"), som opgør den LEVEREDE varme pr. net. Nettet findes af
    kommunens anlægsrækker i EPT (fv_net, også rækker uden produktion); er
    Storkøbenhavns Fjernvarme (net 2) blandt dem, bruges det, ellers nettet med
    størst levering. Dragør, Herlev, Rødovre og Vallensbæk har ingen anlæg i EPT
    og er sat til net 2 i hånden (NET_MANUELT). Kan nettet ikke findes, står
    mixet tomt (status "fælles_net"), og fetch_bolig_fossil.py bruger landssnittet.
    Fossil er kul + olieprodukter + ledningsgas, ikke nettets CO2-faktor, som
    regner biomasse som nul (se metodesiden).

Output:
  data/fjernvarme_mix_scores.csv
  Kolonner: kommune_kode, kommune_navn, fjv_biomasse_pct, fjv_affald_pct,
            fjv_fossil_pct, fjv_ren_pct, fjv_total_tj, fjv_status, fjv_net

Kør fra projektets rodmappe:
  python3 scripts/fetch_fjernvarme_mix.py
"""

import csv
import sys
from collections import defaultdict
from pathlib import Path

import requests
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent))
from kommuner import KOMMUNER  # noqa: E402  (de 98 kommuner, data/kommuner.json)

EPT_URL = "https://ens.dk/media/7199/download"
NET_URL = "https://ens.dk/media/8519/download"
OUTPUT_FIL = Path("data/fjernvarme_mix_scores.csv")
TMP_XLSX = Path("/tmp/ept_fjernvarme.xlsx")
TMP_NET_XLSX = Path("/tmp/ens_fjernvarmenet.xlsx")
STORKBH = 2  # Storkøbenhavns Fjernvarme
# Kommuner på fælles net uden en eneste anlægsrække i EPT (ingen fv_net at slå op).
NET_MANUELT = {"155": STORKBH, "163": STORKBH, "175": STORKBH, "187": STORKBH}
AAR = 2024
MIN_TJ = 100.0  # under dette: egenproduktion er reservekedel/fælles net


BIO = ["skovflis_TJ", "traepiller_TJ", "halm_TJ", "trae- og biomasseaffald_TJ", "bio-olie_TJ", "biogas_TJ"]
AFFALD = ["affald_TJ"]
FOSSIL = ["kul_TJ", "fuelolie_TJ", "spildolie_TJ", "gasolie_TJ", "raffinaderigas_TJ", "lpg_TJ", "naturgas_TJ"]
REN = ["solenergi_TJ", "omgivelsesvarme_TJ", "elektricitet_TJ", "braendselsfrit_TJ"]

# Samme firdeling for net-filen (8519), hvor andelene er brøker 0-1 af leveret varme.
NET_BIO = ["Træ- og biomasseaffald", "Halm", "Skovflis", "Træpiller", "Biogas", "Bioolie"]
NET_AFFALD = ["Affald"]
NET_FOSSIL = ["Kul", "Olieprodukter", "Ledningsgas1"]
NET_REN = ["Overskudsvarme", "Solvarme", "Geotermi", "Omgivelsesvarme", "Elektricitet"]


def download(url=EPT_URL, fil=TMP_XLSX, min_bytes=1_000_000):
    if fil.exists() and fil.stat().st_size > min_bytes:
        print(f"Bruger cached fil: {fil}")
        return
    print(f"Henter {url} ...")
    resp = requests.get(url, timeout=120)
    resp.raise_for_status()
    fil.write_bytes(resp.content)
    print(f"  Gemt: {fil} ({fil.stat().st_size/1024/1024:.1f} MB)")


def net_mix():
    """{net_nr: {navn, tj, bio, affald, fossil, ren}} for AAR fra 8519. Andele i %."""
    wb = openpyxl.load_workbook(TMP_NET_XLSX, read_only=True)
    ws = wb["125% metode"]
    rows = ws.iter_rows(values_only=True)
    next(rows)  # titelrække
    header = list(next(rows))
    idx = {h: i for i, h in enumerate(header) if h}
    mangler = [c for c in ["Fvnet_nr", "År"] + NET_BIO + NET_AFFALD + NET_FOSSIL + NET_REN if c not in idx]
    if mangler:
        raise SystemExit(f"FEJL: ENS har ændret {NET_URL} - mangler kolonner {mangler}")
    tj_kol = next(h for h in idx if str(h).startswith("Varmelevering til net"))

    def g(row, col):
        v = row[idx[col]]
        return float(v) if isinstance(v, (int, float)) else 0.0

    net = {}
    for row in rows:
        nr = row[idx["Fvnet_nr"]]
        if not isinstance(nr, (int, float)) or g(row, "År") != AAR:
            continue
        net[int(nr)] = {
            "navn": row[idx["Fjernvarmenet_navn"]],
            "tj": g(row, tj_kol),
            "bio": 100 * sum(g(row, c) for c in NET_BIO),
            "affald": 100 * sum(g(row, c) for c in NET_AFFALD),
            "fossil": 100 * sum(g(row, c) for c in NET_FOSSIL),
            "ren": 100 * sum(g(row, c) for c in NET_REN),
        }
    # Kontrol: Storkøbenhavns fossilandel var 6,2 % i 2024 (verificeret sep. 2026).
    k = net.get(STORKBH)
    if not k or not (3 <= k["fossil"] <= 12):
        raise SystemExit(f"FEJL: net {STORKBH} ser forkert ud i {NET_URL}: {k}")
    return net


def vaelg_net(kode, net_pr_kommune, net):
    """Nettet der forsyner en kommune uden egen produktion, eller None."""
    if kode in NET_MANUELT:
        return NET_MANUELT[kode]
    kandidater = [n for n in net_pr_kommune.get(kode, ()) if n in net]
    if not kandidater:
        return None
    if STORKBH in kandidater:
        return STORKBH
    return max(kandidater, key=lambda n: net[n]["tj"])


def aggreger():
    wb = openpyxl.load_workbook(TMP_XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    idx = {h: i for i, h in enumerate(header) if h}

    def g(row, col):
        if col not in idx:
            return 0.0
        v = row[idx[col]]
        return float(v) if isinstance(v, (int, float)) else 0.0

    komm = defaultdict(lambda: defaultdict(float))
    net_pr_kommune = defaultdict(set)
    for row in rows:
        if g(row, "aar") != AAR:
            continue
        k = row[idx["vaerk_kommune"]]
        if k is None:
            continue
        k = str(int(k)) if isinstance(k, (int, float)) else str(k).strip()
        if k not in KOMMUNER:  # filtrerer bl.a. Christiansø (411)
            continue
        n = row[idx["fv_net"]] if "fv_net" in idx else None
        if isinstance(n, (int, float)):
            net_pr_kommune[k].add(int(n))  # også anlæg uden produktion (reservekedler)
        if g(row, "varmeprod_TJ") <= 0:
            continue
        a = g(row, "andel_varmelev") or 1.0
        if a > 1:
            a /= 100.0
        komm[k]["bio"] += sum(g(row, c) for c in BIO) * a
        komm[k]["affald"] += sum(g(row, c) for c in AFFALD) * a
        komm[k]["fossil"] += sum(g(row, c) for c in FOSSIL) * a
        komm[k]["ren"] += sum(g(row, c) for c in REN) * a
    return komm, net_pr_kommune


def beregn_og_gem(komm, net_pr_kommune, net):
    resultater = []
    fælles = 0
    for kode, navn in KOMMUNER.items():
        v = komm.get(kode)
        total = (v["bio"] + v["affald"] + v["fossil"] + v["ren"]) if v else 0.0
        if not v or total < MIN_TJ:
            nr = vaelg_net(kode, net_pr_kommune, net)
            m = net.get(nr) if nr is not None else None
            if m:
                resultater.append({
                    "kommune_kode": kode, "kommune_navn": navn,
                    "fjv_biomasse_pct": round(m["bio"], 1), "fjv_affald_pct": round(m["affald"], 1),
                    "fjv_fossil_pct": round(m["fossil"], 1), "fjv_ren_pct": round(m["ren"], 1),
                    "fjv_total_tj": round(total, 1) if v else 0,
                    "fjv_status": "net", "fjv_net": f"{nr} {m['navn']}",
                })
                print(f"  {navn}: net {nr} {m['navn']} (fossil {m['fossil']:.1f}%)")
                continue
            fælles += 1
            print(f"  {navn}: intet net fundet - landssnittet bruges")
            resultater.append({
                "kommune_kode": kode, "kommune_navn": navn,
                "fjv_biomasse_pct": "", "fjv_affald_pct": "",
                "fjv_fossil_pct": "", "fjv_ren_pct": "",
                "fjv_total_tj": round(total, 1) if v else 0,
                "fjv_status": "fælles_net",
            })
            continue
        resultater.append({
            "kommune_kode": kode, "kommune_navn": navn,
            "fjv_biomasse_pct": round(100 * v["bio"] / total, 1),
            "fjv_affald_pct": round(100 * v["affald"] / total, 1),
            "fjv_fossil_pct": round(100 * v["fossil"] / total, 1),
            "fjv_ren_pct": round(100 * v["ren"] / total, 1),
            "fjv_total_tj": round(total, 1),
            "fjv_status": "ok",
        })

    egne = [r for r in resultater if r["fjv_status"] == "ok"]
    if egne:
        nat_bio = sum(komm[r["kommune_kode"]]["bio"] for r in egne)
        nat_tot = sum(r["fjv_total_tj"] for r in egne)
        print(f"  Kommuner med egen varmeproduktion: {len(egne)}/98")
        print(f"  Kommuner med forsynende nets mix: {sum(r['fjv_status'] == 'net' for r in resultater)}/98")
        print(f"  Kommuner uden mix (landssnit): {fælles}/98")
        print(f"  Nationalt biomasse-andel (af egne producenter): {100*nat_bio/nat_tot:.1f}%")
    t = next((r for r in resultater if r["kommune_kode"] == "787"), None)
    if t:
        print(f"  Thisted: biomasse {t['fjv_biomasse_pct']}%  affald {t['fjv_affald_pct']}%  "
              f"fossil {t['fjv_fossil_pct']}%  ren {t['fjv_ren_pct']}%")

    OUTPUT_FIL.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = ["kommune_kode", "kommune_navn", "fjv_biomasse_pct", "fjv_affald_pct",
                  "fjv_fossil_pct", "fjv_ren_pct", "fjv_total_tj", "fjv_status", "fjv_net"]
    with open(OUTPUT_FIL, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        w.writeheader()
        w.writerows(resultater)
    print(f"\n✓ Gemt: {OUTPUT_FIL} ({len(resultater)} kommuner)")


def main():
    print("=" * 65)
    print("Doughnut - Fjernvarmens brændselsmix (Energistyrelsen EPT)")
    print("=" * 65)
    download()
    download(NET_URL, TMP_NET_XLSX, 100_000)
    komm, net_pr_kommune = aggreger()
    beregn_og_gem(komm, net_pr_kommune, net_mix())
    print("\nFærdig! Husk at køre build_master_csv.py.")


if __name__ == "__main__":
    main()
