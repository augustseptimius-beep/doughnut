#!/usr/bin/env python3
"""
fetch_forbrug_co2.py - nutidsjustering af forbrugsbaseret CO2e pr. kommune
==========================================================================

Indikatoren `forbrug_co2` er et Tier 1-estimat (data/methodology_note.md):
kommunernes forbrugsbaserede aftryk i 2011 fra Osei-Owusu m.fl. (2020), ganget
med Energistyrelsens nationale udvikling pr. indbygger fra 2011 til seneste år.
Indtil sep. 2026 var faktoren regnet i hånden (GA25, 2023: 10,11/14,07). Nu
hentes Energistyrelsens tidsserie direkte, så estimatet følger den seneste
Global Afrapportering - også når ENS reviderer de tidligere år, som GA26 gjorde
(2011: 14,07 → 13,46 ton pr. indbygger).

Datakilde: Energistyrelsen, "Danmarks globale klimapåvirkning - Global
afrapportering", datafil bag PowerBI-figurerne "Forbrug", ark "2", rækken
"Udledninger pr. indbygger (ton CO2e)".
  https://ens.dk/analyser-og-statistik/danmarks-globale-klimapaavirkning-global-afrapportering

Faktoren = ENS pr. indbygger (seneste år) / ENS pr. indbygger (2011), begge
fra samme tidsserie, så en metodeforskel mellem Osei-Owusu og ENS ikke
forveksles med en reel reduktion (methodology_note.md afsnit 3.1).

Brug (fra projektets rodmappe):
  python3 scripts/fetch_forbrug_co2.py
  python3 scripts/fetch_forbrug_co2.py --xlsx /sti/til/forbrug.xlsx

Output: data/cba_2023_estimate.csv (kolonnerne cba_estimate, estimat_aar og
skaleringsfaktor; delta_pct er ændringen fra 2011 til estimat_aar. 2011-baselinen
og det gamle håndberegnede 2023-estimat bevares som kildespor).
"""

from __future__ import annotations

import argparse
import csv
import sys
import tempfile
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

ROOT = Path(__file__).resolve().parent.parent
CSV_FIL = ROOT / "data" / "cba_2023_estimate.csv"
GA_URL = "https://ens.dk/media/8503/download"
BASISAAR = 2011
TABEL = "ENS Global Afrapportering"   # nøgle i data/data_years.json


def ens_pr_indbygger(xlsx: Path) -> dict[int, float]:
    """{år: ton CO2e pr. indbygger} fra ENS' datafil."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    if "2" not in wb.sheetnames:
        raise SystemExit(f"FEJL: arket '2' (hovedresultater) findes ikke i {xlsx} - har ENS ændret filen?")
    rows = [[c for c in r if c is not None] for r in wb["2"].iter_rows(values_only=True)]
    aar = next((r for r in rows if r and r[0] == 1990), None)
    pr_indb = next((r for r in rows if r and isinstance(r[0], str) and "pr. indbygger" in r[0]
                    and "CO2e" in r[0]), None)
    if not aar or not pr_indb:
        raise SystemExit("FEJL: fandt ikke rækken 'Udledninger pr. indbygger (ton CO2e)' i ENS' datafil")
    return {int(a): float(v) for a, v in zip(aar, pr_indb[1:])}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="Lokal kopi af ENS' datafil (hentes ellers)")
    args = ap.parse_args()
    xlsx = Path(args.xlsx) if args.xlsx else Path(tempfile.gettempdir()) / "ens_ga_forbrug.xlsx"
    if not args.xlsx:
        print(f"Henter ENS' datafil: {GA_URL}")
        req = urllib.request.Request(GA_URL, headers={"User-Agent": "Mozilla/5.0 DoughnutDK/1.0"})
        with urllib.request.urlopen(req, timeout=120) as r, open(xlsx, "wb") as f:
            f.write(r.read())
    serie = ens_pr_indbygger(xlsx)
    sidste = max(serie)
    faktor = serie[sidste] / serie[BASISAAR]
    print(f"  ENS pr. indbygger: {BASISAAR} = {serie[BASISAAR]:.2f} t, {sidste} = {serie[sidste]:.2f} t")
    print(f"  Skaleringsfaktor {BASISAAR}→{sidste}: {faktor:.4f}")

    with open(CSV_FIL, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r["cba_estimate"] = round(float(r["cba_2011_baseline"]) * faktor, 2)
        r["estimat_aar"] = sidste
        r["skaleringsfaktor"] = round(faktor, 4)
        r["delta_pct"] = round((faktor - 1) * 100, 1)   # ændring 2011 → estimat_aar
    felter = ["kommune_kode", "kommune", "cba_2011_baseline", "cba_2023_estimate", "cba_estimate",
              "estimat_aar", "skaleringsfaktor", "delta_pct", "confidence_flag"]
    with open(CSV_FIL, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=felter, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)
    print(f"✓ {CSV_FIL.relative_to(ROOT)}: {len(rows)} kommuner, estimat for {sidste}")

    from dst_aar import registrer_aar
    registrer_aar(TABEL, str(sidste))
    t = next((r for r in rows if r["kommune_kode"] == "787"), None)
    if t:
        print(f"  Thisted: {t['cba_2011_baseline']} → {t['cba_estimate']} t CO2e/indb.")
    return 0


if __name__ == "__main__":
    rc = main()
    try:
        from build_master_csv import auto_build_master
        auto_build_master()
    except Exception as _e:
        print(f"\n⚠ Kunne ikke auto-rebuild master-CSV: {_e}")
        print("  Rådata er gemt. Kør manuelt: python3 scripts/build_master_csv.py")
    sys.exit(rc)
